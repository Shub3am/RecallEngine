import csv
import json
import os
import sqlite3
from pathlib import Path

import pytest

from recall_engine.sources import (
    PASSAGE_ID_KEY,
    SOURCE_KEY,
    SUPPORTED_FILE_SUFFIXES,
    load_documents,
    source_fingerprint,
)
from recall_engine.sources.chunking import chunk_text


def write_minimal_pdf(path: Path, page_texts: list[str]) -> None:
    """Hand-written PDF with one Helvetica text line per page, because pypdf cannot author text."""
    page_count = len(page_texts)
    font_object_number = 3 + 2 * page_count
    page_object_numbers = [3 + 2 * index for index in range(page_count)]
    object_bodies = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [%s] /Count %d >>"
        % (b" ".join(b"%d 0 R" % number for number in page_object_numbers), page_count),
    ]
    for page_object_number, page_text in zip(page_object_numbers, page_texts):
        content_stream = b"BT /F1 12 Tf 72 720 Td (%s) Tj ET" % page_text.encode("latin-1")
        object_bodies.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R "
            b"/Resources << /Font << /F1 %d 0 R >> >> >>" % (page_object_number + 1, font_object_number)
        )
        object_bodies.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content_stream), content_stream))
    object_bodies.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    pdf_bytes = bytearray(b"%PDF-1.4\n")
    object_offsets = []
    for object_number, object_body in enumerate(object_bodies, start=1):
        object_offsets.append(len(pdf_bytes))
        pdf_bytes += b"%d 0 obj\n%s\nendobj\n" % (object_number, object_body)
    xref_offset = len(pdf_bytes)
    pdf_bytes += b"xref\n0 %d\n0000000000 65535 f \n" % (len(object_bodies) + 1)
    for object_offset in object_offsets:
        pdf_bytes += b"%010d 00000 n \n" % object_offset
    pdf_bytes += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(object_bodies) + 1,
        xref_offset,
    )
    path.write_bytes(bytes(pdf_bytes))


def numbered_words(count: int) -> str:
    return " ".join(f"word{number}" for number in range(count))


@pytest.fixture
def sqlite_path(tmp_path: Path) -> Path:
    database_path = tmp_path / "library.db"
    with sqlite3.connect(database_path) as connection:
        connection.execute("CREATE TABLE books (id INTEGER, title TEXT, price NUMERIC, added DATE)")
        connection.executemany(
            "INSERT INTO books VALUES (?, ?, ?, ?)",
            [(1, "Red Apple", 9.5, "2024-01-01"), (2, "Green Banana", 12, "2024-02-01")],
        )
        connection.execute("CREATE TABLE authors (name TEXT)")
        connection.execute("INSERT INTO authors VALUES ('Ada')")
    connection.close()
    return database_path


def test_chunk_text_windows_overlap_by_forty_words() -> None:
    chunks = chunk_text(numbered_words(450))

    assert len(chunks) == 3
    assert [len(chunk.split()) for chunk in chunks] == [200, 200, 130]
    assert chunks[0].split()[-40:] == chunks[1].split()[:40]
    assert chunks[1].split()[0] == "word160"


def test_chunk_text_handles_short_and_blank_text() -> None:
    assert chunk_text("  \n\t ") == []
    assert chunk_text("just a few words") == ["just a few words"]
    assert len(chunk_text(numbered_words(200))) == 1


def test_json_list_keeps_fields_and_skips_non_dicts(tmp_path: Path) -> None:
    json_path = tmp_path / "items.json"
    json_path.write_text(json.dumps([{"id": 7, "title": "A", "tags": ["x"]}, "noise", {"id": 8, "title": "B"}]))

    documents = load_documents(str(json_path))

    assert documents == [
        {"id": 7, "title": "A", "tags": ["x"], PASSAGE_ID_KEY: "items.json#1", SOURCE_KEY: "items.json"},
        {"id": 8, "title": "B", PASSAGE_ID_KEY: "items.json#2", SOURCE_KEY: "items.json"},
    ]


def test_json_dict_with_data_key(tmp_path: Path) -> None:
    json_path = tmp_path / "movies.json"
    json_path.write_text(json.dumps({"movies": [{"title": "A"}], "genres": [{"name": "drama"}]}))

    documents = load_documents(str(json_path), data_key="genres")

    assert [document["name"] for document in documents] == ["drama"]
    with pytest.raises(ValueError, match="missing"):
        load_documents(str(json_path), data_key="missing")


def test_json_dict_auto_detects_single_list(tmp_path: Path) -> None:
    single_list_path = tmp_path / "single.json"
    single_list_path.write_text(json.dumps({"total": 2, "movies": [{"title": "A"}, {"title": "B"}]}))
    two_lists_path = tmp_path / "two.json"
    two_lists_path.write_text(json.dumps({"a": [1], "b": [2], "name": "whole"}))

    assert [document["title"] for document in load_documents(str(single_list_path))] == ["A", "B"]
    assert load_documents(str(two_lists_path)) == [
        {"a": [1], "b": [2], "name": "whole", PASSAGE_ID_KEY: "two.json#1", SOURCE_KEY: "two.json"}
    ]


def test_generated_keys_overwrite_existing_fields(tmp_path: Path) -> None:
    json_path = tmp_path / "clash.json"
    json_path.write_text(json.dumps([{"passage_id": "mine", "source": "elsewhere", "title": "A"}]))

    documents = load_documents(str(json_path))

    assert documents[0][PASSAGE_ID_KEY] == "clash.json#1"
    assert documents[0][SOURCE_KEY] == "clash.json"


def test_json_lines_skip_blank_lines(tmp_path: Path) -> None:
    json_lines_path = tmp_path / "events.jsonl"
    json_lines_path.write_text('{"event": "start"}\n\n{"event": "stop"}\n')

    documents = load_documents(str(json_lines_path))

    assert [document["event"] for document in documents] == ["start", "stop"]
    assert documents[1][PASSAGE_ID_KEY] == "events.jsonl#2"


@pytest.mark.parametrize(("file_name", "delimiter"), [("people.csv", ","), ("people.tsv", "\t")])
def test_delimited_files(tmp_path: Path, file_name: str, delimiter: str) -> None:
    delimited_path = tmp_path / file_name
    with delimited_path.open("w", encoding="utf-8", newline="") as delimited_file:
        writer = csv.writer(delimited_file, delimiter=delimiter)
        writer.writerow(["name", "city"])
        writer.writerow(["Ada", "London, UK"])
        writer.writerow(["Alan", "Wilmslow"])

    documents = load_documents(str(delimited_path))

    assert documents[0] == {"name": "Ada", "city": "London, UK", PASSAGE_ID_KEY: f"{file_name}#1", SOURCE_KEY: file_name}
    assert documents[1]["city"] == "Wilmslow"


def test_long_text_file_is_chunked_with_overlap(tmp_path: Path) -> None:
    text_path = tmp_path / "essay.txt"
    text_path.write_text(numbered_words(450), encoding="utf-8")

    documents = load_documents(str(text_path))

    assert [document[PASSAGE_ID_KEY] for document in documents] == ["essay.txt#1", "essay.txt#2", "essay.txt#3"]
    assert documents[0]["text"].split()[-40:] == documents[1]["text"].split()[:40]


def test_markdown_file(tmp_path: Path) -> None:
    markdown_path = tmp_path / "notes.md"
    markdown_path.write_text("# Title\n\nSome *markdown* body.\n", encoding="utf-8")

    assert load_documents(str(markdown_path)) == [
        {"text": "# Title Some *markdown* body.", PASSAGE_ID_KEY: "notes.md#1", SOURCE_KEY: "notes.md"}
    ]


def test_html_skips_script_style_and_head(tmp_path: Path) -> None:
    html_path = tmp_path / "page.html"
    html_path.write_text(
        "<html><head><title>Hidden title</title><style>p {color: red}</style></head>"
        "<body><h1>Visible   heading</h1><script>var secret = 1;</script><p>Body\n text</p></body></html>",
        encoding="utf-8",
    )

    documents = load_documents(str(html_path))

    assert [document["text"] for document in documents] == ["Visible heading Body text"]


def test_pdf_pages_are_numbered_from_one(tmp_path: Path) -> None:
    pdf_path = tmp_path / "report.pdf"
    write_minimal_pdf(pdf_path, ["First page words", "Second page words"])

    documents = load_documents(str(pdf_path))

    assert [(document["text"], document["page"]) for document in documents] == [
        ("First page words", 1),
        ("Second page words", 2),
    ]
    assert documents[1][PASSAGE_ID_KEY] == "report.pdf#2"


def test_docx_paragraphs_are_joined(tmp_path: Path) -> None:
    import docx

    docx_path = tmp_path / "memo.docx"
    memo = docx.Document()
    memo.add_paragraph("Quarterly memo")
    memo.add_paragraph("Revenue grew.")
    memo.save(str(docx_path))

    assert load_documents(str(docx_path)) == [
        {"text": "Quarterly memo Revenue grew.", PASSAGE_ID_KEY: "memo.docx#1", SOURCE_KEY: "memo.docx"}
    ]


def test_xlsx_sheets_continue_numbering(tmp_path: Path) -> None:
    import datetime

    import openpyxl

    xlsx_path = tmp_path / "sales.xlsx"
    workbook = openpyxl.Workbook()
    first_sheet = workbook.active
    first_sheet.title = "Q1"
    first_sheet.append(["region", None, "amount"])
    first_sheet.append(["north", "ignored", 10])
    first_sheet.append([None, None, None])
    first_sheet.append(["south", "ignored", 20])
    second_sheet = workbook.create_sheet("Q2")
    second_sheet.append(["region", "closed"])
    second_sheet.append(["east", datetime.date(2024, 4, 1)])
    workbook.save(xlsx_path)

    documents = load_documents(str(xlsx_path))

    assert documents == [
        {"region": "north", "amount": 10, "sheet": "Q1", PASSAGE_ID_KEY: "sales.xlsx#1", SOURCE_KEY: "sales.xlsx"},
        {"region": "south", "amount": 20, "sheet": "Q1", PASSAGE_ID_KEY: "sales.xlsx#2", SOURCE_KEY: "sales.xlsx"},
        {
            "region": "east",
            "closed": "2024-04-01 00:00:00",
            "sheet": "Q2",
            PASSAGE_ID_KEY: "sales.xlsx#3",
            SOURCE_KEY: "sales.xlsx",
        },
    ]
    json.dumps(documents)


def test_sqlite_loads_every_table(sqlite_path: Path) -> None:
    documents = load_documents(str(sqlite_path))

    assert [document[PASSAGE_ID_KEY] for document in documents] == ["authors#1", "books#1", "books#2"]
    assert documents[1] == {
        "id": 1,
        "title": "Red Apple",
        "price": 9.5,
        "added": "2024-01-01",
        PASSAGE_ID_KEY: "books#1",
        SOURCE_KEY: "books",
    }


def test_sqlite_table_and_query(sqlite_path: Path) -> None:
    table_documents = load_documents(str(sqlite_path), table="authors")
    query_documents = load_documents(
        str(sqlite_path), table="authors", query="SELECT title FROM books WHERE price > 10"
    )

    assert table_documents == [{"name": "Ada", PASSAGE_ID_KEY: "authors#1", SOURCE_KEY: "authors"}]
    assert query_documents == [{"title": "Green Banana", PASSAGE_ID_KEY: "query#1", SOURCE_KEY: "query"}]


def test_sqlalchemy_url(sqlite_path: Path) -> None:
    database_url = f"sqlite:///{sqlite_path}"

    all_documents = load_documents(database_url)
    table_documents = load_documents(database_url, table="books")
    query_documents = load_documents(database_url, query="SELECT name FROM authors")

    assert sorted(document[PASSAGE_ID_KEY] for document in all_documents) == ["authors#1", "books#1", "books#2"]
    assert [document["title"] for document in table_documents] == ["Red Apple", "Green Banana"]
    assert query_documents == [{"name": "Ada", PASSAGE_ID_KEY: "query#1", SOURCE_KEY: "query"}]
    json.dumps(all_documents)


def build_folder(root: Path) -> Path:
    (root / "nested" / "deeper").mkdir(parents=True)
    (root / ".hidden_dir").mkdir()
    (root / "b.txt").write_text("bravo text", encoding="utf-8")
    (root / "a.json").write_text(json.dumps([{"title": "alpha"}]), encoding="utf-8")
    (root / "nested" / "c.md").write_text("charlie text", encoding="utf-8")
    (root / "nested" / "deeper" / "d.csv").write_text("name\ndelta\n", encoding="utf-8")
    (root / ".secret.txt").write_text("hidden file", encoding="utf-8")
    (root / ".hidden_dir" / "e.txt").write_text("hidden dir", encoding="utf-8")
    (root / "blob.bin").write_bytes(b"\x00\x01")
    return root


def test_folder_walk_is_recursive_and_deterministic(tmp_path: Path) -> None:
    folder = build_folder(tmp_path / "corpus")

    documents = load_documents(str(folder))

    assert [(document[PASSAGE_ID_KEY], document[SOURCE_KEY]) for document in documents] == [
        ("a.json#1", "a.json"),
        ("b.txt#1", "b.txt"),
        ("nested/c.md#1", "nested/c.md"),
        ("nested/deeper/d.csv#1", "nested/deeper/d.csv"),
    ]
    assert load_documents(str(folder)) == documents


def test_unsupported_suffix_and_missing_path(tmp_path: Path) -> None:
    binary_path = tmp_path / "blob.bin"
    binary_path.write_bytes(b"\x00")

    with pytest.raises(ValueError, match=r"\.pdf"):
        load_documents(str(binary_path))
    with pytest.raises(FileNotFoundError):
        load_documents(str(tmp_path / "missing.json"))
    assert ".xlsx" in SUPPORTED_FILE_SUFFIXES


def test_file_fingerprint_is_stable_and_tracks_data_key(tmp_path: Path) -> None:
    json_path = tmp_path / "items.json"
    json_path.write_text("[]", encoding="utf-8")

    fingerprint = source_fingerprint(str(json_path))

    assert fingerprint == source_fingerprint(str(json_path))
    assert fingerprint["kind"] == "file"
    assert fingerprint["chunk_words"] == 200
    assert source_fingerprint(str(json_path), data_key="movies") != fingerprint


def test_folder_fingerprint_changes_when_file_modified_or_added(tmp_path: Path) -> None:
    folder = build_folder(tmp_path / "corpus")
    original_fingerprint = source_fingerprint(str(folder))
    assert source_fingerprint(str(folder)) == original_fingerprint
    assert [entry[0] for entry in original_fingerprint["files"]] == [
        "a.json",
        "b.txt",
        "nested/c.md",
        "nested/deeper/d.csv",
    ]

    text_path = folder / "b.txt"
    text_stat = text_path.stat()
    os.utime(text_path, ns=(text_stat.st_atime_ns, text_stat.st_mtime_ns + 1_000_000_000))
    modified_fingerprint = source_fingerprint(str(folder))
    assert modified_fingerprint != original_fingerprint

    (folder / "nested" / "new.txt").write_text("new file", encoding="utf-8")
    assert source_fingerprint(str(folder)) != modified_fingerprint


def test_sqlite_fingerprint_and_url_fingerprint(sqlite_path: Path) -> None:
    fingerprint = source_fingerprint(str(sqlite_path), table="books")

    assert fingerprint["kind"] == "sqlite"
    assert fingerprint["table"] == "books"
    assert source_fingerprint(str(sqlite_path), query="SELECT 1") != fingerprint
    assert source_fingerprint(f"sqlite:///{sqlite_path}") is None
